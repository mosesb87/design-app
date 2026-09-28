"""Builds Review-Engine.xlsx (the master workbook) and one QA workbook per target.

Run from this folder:  python3 build_workbook.py
Outputs:
  Review-Engine.xlsx                                  — Rubric · Index (21 + targets, live formulas) · Checks · Findings library ·
                                                        Impact ledger · Outreach · Messages
  targets/<slug>/site/files/<Company>-…-Workbook.xlsx — the target's own scorecard, findings and weekly QA sheet
Recalculate with LibreOffice after building so cached values exist:  soffice --headless --convert-to xlsx --outdir tmp Review-Engine.xlsx
"""
import json, os, re, datetime
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import CellIsRule, DataBarRule
from engine import (HERE, PILLARS, CHECKS, CHECK_IDS, WEIGHTS, GRADES, BANDS, load_index, load_targets,
                    rank_against_index, weakest_pillar)

TODAY = "2026-09-28"
index = load_index()
targets = load_targets()
ALL = index + targets                      # the scoring pool: benchmark stores first, then targets
N = len(ALL); LAST = N + 1

# ---------- styles (same language as the ShelfMark workbook) ----------
ARIAL = "Arial"
INK = "1F1D1A"; PAPER = "F4F1EA"; ACCENT = "C8451F"; MUTED = "6B6660"; LINE = "D9D4CA"; COBALT = "1F3FBF"
f_base = Font(name=ARIAL, size=10, color=INK)
f_bold = Font(name=ARIAL, size=10, bold=True, color=INK)
f_head = Font(name=ARIAL, size=10, bold=True, color="FFFFFF")
f_title = Font(name=ARIAL, size=18, bold=True, color=INK)
f_sub = Font(name=ARIAL, size=11, color=MUTED)
f_small = Font(name=ARIAL, size=8, color=MUTED)
f_input = Font(name=ARIAL, size=10, color="0000FF")
f_link = Font(name=ARIAL, size=10, color="0000FF", underline="single")
f_green = Font(name=ARIAL, size=10, color="008000")
fill_head = PatternFill("solid", fgColor=INK)
fill_target = PatternFill("solid", fgColor="E8EEFF")
fill_input = PatternFill("solid", fgColor="FFFF00")
fill_soft = PatternFill("solid", fgColor=PAPER)
thin = Side(style="thin", color=LINE)
box = Border(left=thin, right=thin, top=thin, bottom=thin)
wrap = Alignment(wrap_text=True, vertical="top")
center = Alignment(horizontal="center", vertical="center")


def style_header(ws, row, ncols):
    for c in range(1, ncols + 1):
        cell = ws.cell(row=row, column=c)
        cell.fill = fill_head; cell.font = f_head; cell.border = box
        cell.alignment = Alignment(wrap_text=True, vertical="center")


def set_widths(ws, widths):
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w


def finish(ws, r, ncols, wrap_cols=(), center_cols=()):
    for col in range(1, ncols + 1):
        c = ws.cell(row=r, column=col); c.border = box
        if c.font == Font():
            c.font = f_base
        c.alignment = Alignment(vertical="top", wrap_text=col in wrap_cols, horizontal="center" if col in center_cols else None)


def readme_block(ws, title, subtitle, lines, start=2):
    ws.sheet_view.showGridLines = False
    set_widths(ws, [3, 26, 96])
    ws.cell(row=start, column=2, value=title).font = f_title
    ws.cell(row=start + 1, column=2, value=subtitle).font = f_sub
    r = start + 3
    for k, v in lines:
        ws.cell(row=r, column=2, value=k).font = f_bold
        ws.cell(row=r, column=2).alignment = wrap
        c = ws.cell(row=r, column=3, value=v); c.font = f_base; c.alignment = wrap
        ws.row_dimensions[r].height = max(30, 15 * (len(v) // 100 + 1))
        r += 1
    return r


# References that ground the rubric (published, checkable). Only checks with a real source get one.
REFERENCES = {
    "F2": "Google Search Central — robots.txt introduction and the Sitemap directive · https://developers.google.com/search/docs/crawling-indexing/robots/intro",
    "F3": "Google Search Central — Build and submit a sitemap (50,000 URLs / 50 MB per file) · https://developers.google.com/search/docs/crawling-indexing/sitemaps/build-sitemap",
    "F4": "Google Search Central — Influencing your title links in search results · https://developers.google.com/search/docs/appearance/title-link",
    "F5": "Google Search Central — Control your snippets in search results (unique descriptions per page) · https://developers.google.com/search/docs/appearance/snippet",
    "F6": "Google Search Central — Consolidate duplicate URLs with canonicals · https://developers.google.com/search/docs/crawling-indexing/consolidate-duplicate-urls",
    "F7": "Google Search Central — Product structured data (Offer, availability, aggregateRating rules) · https://developers.google.com/search/docs/appearance/structured-data/product",
    "F8": "Google Search Central — URL structure best practices · https://developers.google.com/search/docs/crawling-indexing/url-structure",
    "B2": "FTC Mail, Internet, or Telephone Order Merchandise Rule (shipping-time representations) · https://www.ftc.gov/legal-library/browse/rules/mail-internet-or-telephone-order-merchandise-rule",
    "T5": "WCAG 2.2 — 1.1.1 Non-text Content and 2.4.1 Bypass Blocks · https://www.w3.org/TR/WCAG22/",
    "N1": "Google Search Central — Creating helpful, reliable, people-first content · https://developers.google.com/search/docs/fundamentals/creating-helpful-content",
}
SPEED_REF = "Deloitte Digital for Google, “Milliseconds Make Millions” (2020): a 0.1 s mobile speed improvement was associated with +8.4% retail conversion and +9.2% AOV (correlational). Core Web Vitals thresholds (web.dev): LCP ≤ 2.5 s, INP ≤ 200 ms, CLS ≤ 0.1. Used only as context — the rubric does not score speed."


def build_master():
    wb = Workbook()
    # ---------------- README ----------------
    ws = wb.active; ws.title = "README"
    r = readme_block(ws, "Review Engine", f"ShelfMark's 29 checks as a standing rubric for any store · index of {len(index)} benchmark stores + {len(targets)} reviewed targets · {TODAY} · Mousa Batarseh", [
        ("What this is", "The spreadsheet behind every store review I send with an application. One rubric (29 public-page checks, 5 pillars, 0–100 score) scores any storefront in about an hour; the ShelfMark index of 21 Michigan stores is the comparison pool, so each target gets a rank, not just a number; and every role-specific finding is filed in a library that the next review reuses."),
        ("How a score is built", "Each check is 2 (meets), 1 (partly), 0 (fails) or U (unverified from public pages). Pillar % = verified points ÷ (2 × verified checks). Score = weighted average of pillar % (weights on Rubric: Findability 25, Catalog 30, Buy 20, Trust 15, Content 10). U never counts for or against a store; Coverage and Confidence show how much was observable. Identical to the ShelfMark workbook, so scores are comparable."),
        ("Sheets", "Rubric — pillars, weights, thresholds, the 29 checks and the published reference behind each one. Index — every store scored, benchmark and targets together (targets shaded blue), with live formulas. Checks — 29 observations per store with dated evidence; the only place scores are typed. Findings — the library: strengths, fixes and role-specific findings with counts, the duty they map to, the fix and the evidence file. Impact ledger — my own outcomes with source and confidence, the single source of truth for résumé bullets. Outreach — one row per person contacted, with status and follow-up dates. Messages — the notes that go with a review."),
        ("The one-hour review", "1) Paste the target's homepage, robots.txt, sitemap, one category and two product URLs into the audit brief (audit-brief.md). 2) Score the 29 checks with evidence into targets/<slug>/audit.json. 3) Read the posting; find 4–6 findings that can be COUNTED across sampled pages and map each to a duty in the posting. 4) python3 build_workbook.py && python3 build_review.py <slug> — the scorecard, rank, workbook and review page are generated. 5) Write the working fixes (prototypes) by hand in the company's look. 6) Send: application link + one finding to one named person within 48 hours; log the row on Outreach."),
        ("Rules", "Public pages only; never log in, cart or submit forms. Count things (14 of 20), never characterize (\"many\"). Lead with three strengths. Every finding names the page, the count, the duty it maps to and the fix. No claims about health, legal or financial language. Review pages are noindex and sent as private links to one person."),
        ("Legend", "Blue text on yellow = typed inputs (scores, contacts, statuses). Black = formulas. Green = pulled from another sheet. Blue-shaded rows on Index = reviewed targets (not part of the public benchmark)."),
        ("Author", "Mousa Batarseh · hireme@mousabatarseh.com · mousabatarseh.com · linkedin.com/in/mousabatarseh"),
    ])

    # ---------------- Rubric ----------------
    wr = wb.create_sheet("Rubric")
    wr.sheet_view.showGridLines = False
    set_widths(wr, [8, 24, 10, 56, 44, 44, 44, 70])
    wr["A1"] = "Pillars & weights"; wr["A1"].font = f_title
    for i, h in enumerate(["Pillar", "Name", "Weight", "Question the pillar answers"], 1):
        wr.cell(row=2, column=i, value=h)
    style_header(wr, 2, 4)
    for i, (pid, name, w, q) in enumerate(PILLARS, start=3):
        wr.cell(row=i, column=1, value=pid); wr.cell(row=i, column=2, value=name)
        c = wr.cell(row=i, column=3, value=w); c.font = f_input; c.fill = fill_input
        wr.cell(row=i, column=4, value=q)
        finish(wr, i, 4, wrap_cols=(4,))
    wr.cell(row=8, column=2, value="Total").font = f_bold
    wr.cell(row=8, column=3, value="=SUM(C3:C7)").font = f_bold
    wr["A10"] = "Grade thresholds (score ≥)"; wr["A10"].font = f_bold
    for i, (cut, g) in enumerate(GRADES[:4], start=11):
        wr.cell(row=i, column=1, value=g); wr.cell(row=i, column=2, value=BANDS[g])
        c = wr.cell(row=i, column=3, value=cut); c.font = f_input; c.fill = fill_input
        finish(wr, i, 3)
    wr.cell(row=15, column=1, value="F"); wr.cell(row=15, column=2, value=BANDS["F"]); wr.cell(row=15, column=3, value="below D")
    finish(wr, 15, 3)
    wr["A16"] = "Confidence: High = 26+ verified checks, Medium = 20–25, Low = under 20 (of 29)."; wr["A16"].font = f_sub
    wr["A17"] = SPEED_REF; wr["A17"].font = f_small; wr["A17"].alignment = wrap
    wr.merge_cells("A17:H17"); wr.row_dimensions[17].height = 30
    wr["A19"] = "The 29 checks"; wr["A19"].font = f_title
    for i, h in enumerate(["ID", "Check", "Pillar", "Scores 2 when…", "Scores 1 when…", "Scores 0 when…", "Published reference (where one exists)"], 1):
        wr.cell(row=20, column=i if i < 7 else 8, value=h)
    wr.cell(row=20, column=7, value="")
    style_header(wr, 20, 8)
    for i, (cid, name, s2, s1, s0) in enumerate(CHECKS, start=21):
        for col, v in enumerate([cid, name, cid[0], s2, s1, s0, "", REFERENCES.get(cid, "")], 1):
            wr.cell(row=i, column=col, value=v)
        finish(wr, i, 8, wrap_cols=(4, 5, 6, 8))
        wr.row_dimensions[i].height = 44
    wr.freeze_panes = "A21"

    # ---------------- Checks (long) ----------------
    wc = wb.create_sheet("Checks")
    set_widths(wc, [34, 9, 8, 9, 24, 8, 110])
    hdr = ["Store", "Source", "Pillar", "Check ID", "Check", "Score", f"Evidence (what was observed, {TODAY})"]
    for i, h in enumerate(hdr, 1):
        wc.cell(row=1, column=i, value=h)
    style_header(wc, 1, len(hdr))
    cr = 2
    for st in ALL:
        for cid, name, *_ in CHECKS:
            chk = st["checks"][cid]
            for col, v in enumerate([st["company"], "Benchmark" if st["source"] == "index" else "Target", cid[0], cid, name, chk["s"], chk["e"]], 1):
                wc.cell(row=cr, column=col, value=v)
            finish(wc, cr, 7, wrap_cols=(7,), center_cols=(3, 4, 6))
            sc = wc.cell(row=cr, column=6); sc.font = f_input; sc.fill = fill_input
            if st["source"] == "target":
                wc.cell(row=cr, column=1).fill = fill_target
            cr += 1
    CHK_LAST = cr - 1
    dv = DataValidation(type="list", formula1='"0,1,2,U"', allow_blank=False)
    wc.add_data_validation(dv); dv.add(f"F2:F{CHK_LAST + 600}")
    wc.freeze_panes = "A2"; wc.auto_filter.ref = f"A1:G{CHK_LAST}"

    # ---------------- Index ----------------
    wi = wb.create_sheet("Index", 1)
    wi.sheet_view.showGridLines = False
    cols = ["Rank (all)", "Rank on the benchmark", "Store", "Source", "Segment", "Platform", "HQ", "URL", "Audited",
            "Findability %", "F n", "Catalog %", "C n", "Buy %", "B n", "Trust %", "T n", "Content %", "N n",
            "Verified (of 29)", "Coverage", "Confidence", "Score", "Grade", "Band", "Weakest pillar", "Review page", "Notes"]
    for i, h in enumerate(cols, 1):
        wi.cell(row=1, column=i, value=h)
    style_header(wi, 1, len(cols)); wi.row_dimensions[1].height = 34
    set_widths(wi, [7, 10, 34, 10, 18, 22, 20, 34, 11, 11, 4, 11, 4, 9, 4, 9, 4, 11, 4, 9, 9, 10, 9, 7, 12, 10, 40, 50])
    W = {p: f"Rubric!$C${3 + i}" for i, (p, *_) in enumerate(PILLARS)}
    PC = {"F": ("J", "K"), "C": ("L", "M"), "B": ("N", "O"), "T": ("P", "Q"), "N": ("R", "S")}
    RNG = CHK_LAST + 600
    CHK = f"Checks!$A$2:$A${RNG}"; CHP = f"Checks!$C$2:$C${RNG}"; CHS = f"Checks!$F$2:$F${RNG}"
    for idx, st in enumerate(ALL):
        r = idx + 2
        wi.cell(row=r, column=1, value=f'=IF(W{r}="","",RANK(W{r},$W$2:$W${LAST},0))')
        # rank among benchmark stores only (targets are ranked "as if" they joined): 1 + count of benchmark scores above
        wi.cell(row=r, column=2, value=f'=IF(W{r}="","",1+COUNTIFS($D$2:$D${LAST},"Benchmark",$W$2:$W${LAST},">"&W{r}))')
        wi.cell(row=r, column=3, value=st["company"]).font = f_bold
        wi.cell(row=r, column=4, value="Benchmark" if st["source"] == "index" else "Target")
        wi.cell(row=r, column=5, value=st["segment"]); wi.cell(row=r, column=6, value=st["platform"]); wi.cell(row=r, column=7, value=st["hq_city"])
        u = wi.cell(row=r, column=8, value=st["url"]); u.hyperlink = st["url"]; u.font = f_link
        wi.cell(row=r, column=9, value=st["audited_at"])
        for p, (pcol, ncol) in PC.items():
            wi[f"{pcol}{r}"] = (f'=IFERROR(SUMIFS({CHS},{CHK},$C{r},{CHP},"{p}")/(2*COUNTIFS({CHK},$C{r},{CHP},"{p}",{CHS},">=0")),"")')
            wi[f"{pcol}{r}"].number_format = "0%"
            wi[f"{ncol}{r}"] = f'=COUNTIFS({CHK},$C{r},{CHP},"{p}",{CHS},">=0")'; wi[f"{ncol}{r}"].font = f_small
        wi[f"T{r}"] = f"=K{r}+M{r}+O{r}+Q{r}+S{r}"
        wi[f"U{r}"] = f"=T{r}/29"; wi[f"U{r}"].number_format = "0%"
        wi[f"V{r}"] = f'=IF(T{r}>=26,"High",IF(T{r}>=20,"Medium","Low"))'
        num = "+".join(f'IF({PC[p][1]}{r}>0,{W[p]}*{PC[p][0]}{r},0)' for p in PC)
        den = "+".join(f'IF({PC[p][1]}{r}>0,{W[p]},0)' for p in PC)
        wi[f"W{r}"] = f'=IF(({den})=0,"",ROUND(100*({num})/({den}),1))'
        wi[f"W{r}"].font = Font(name=ARIAL, size=11, bold=True, color=INK); wi[f"W{r}"].number_format = "0.0"
        wi[f"X{r}"] = (f'=IF(W{r}="","",IF(W{r}>=Rubric!$C$11,"A",IF(W{r}>=Rubric!$C$12,"B",IF(W{r}>=Rubric!$C$13,"C",IF(W{r}>=Rubric!$C$14,"D","F")))))')
        wi[f"Y{r}"] = f'=IF(X{r}="","",IF(X{r}="A","Benchmark",IF(X{r}="B","Strong",IF(X{r}="C","Developing",IF(X{r}="D","Needs work","At risk")))))'
        wi[f"Z{r}"] = (f'=IF(COUNT(J{r},L{r},N{r},P{r},R{r})=0,"",INDEX({{"Findability","Catalog","Buy","Trust","Content"}},MATCH(MIN(J{r},L{r},N{r},P{r},R{r}),CHOOSE({{1,2,3,4,5}},J{r},L{r},N{r},P{r},R{r}),0)))')
        rv = st.get("review")
        if rv:
            link = f"https://mousabatarseh.com/{rv['path']}/"
            c = wi.cell(row=r, column=27, value=link); c.hyperlink = link; c.font = f_link
        wi.cell(row=r, column=28, value=st.get("notes", ""))
        finish(wi, r, len(cols), wrap_cols=(28,), center_cols=(1, 2, 4, 9, 11, 13, 15, 17, 19, 20, 21, 22, 24))
        if st["source"] == "target":
            for col in range(1, len(cols) + 1):
                wi.cell(row=r, column=col).fill = fill_target
        wi.row_dimensions[r].height = 30
    wi.freeze_panes = "D2"; wi.auto_filter.ref = f"A1:AB{LAST}"
    wi.conditional_formatting.add(f"W2:W{LAST}", DataBarRule(start_type="num", start_value=0, end_type="num", end_value=100, color=ACCENT))
    for g, colr in (("A", "C6EFCE"), ("B", "E2F0D9"), ("C", "FFF2CC"), ("D", "FCE4D6"), ("F", "F8CBAD")):
        wi.conditional_formatting.add(f"X2:X{LAST}", CellIsRule(operator="equal", formula=[f'"{g}"'], fill=PatternFill("solid", fgColor=colr)))
    wi.cell(row=LAST + 2, column=3, value="Benchmark average").font = f_bold
    wi.cell(row=LAST + 2, column=23, value=f'=ROUND(AVERAGEIF(D2:D{LAST},"Benchmark",W2:W{LAST}),1)').font = f_bold
    wi.cell(row=LAST + 3, column=3, value="Benchmark median").font = f_bold
    wi.cell(row=LAST + 4, column=3, value="Rank on the benchmark = where a target would place among the 21 Michigan stores if it joined the index (1 = best). Targets are shaded blue and are not part of the public ShelfMark index.").font = f_sub
    wi.merge_cells(start_row=LAST + 4, start_column=3, end_row=LAST + 4, end_column=14); wi.cell(row=LAST + 4, column=3).alignment = wrap; wi.row_dimensions[LAST + 4].height = 30

    # ---------------- Findings library ----------------
    wf = wb.create_sheet("Findings")
    set_widths(wf, [30, 10, 8, 11, 7, 60, 44, 40, 60, 60, 44, 12, 14])
    hdr = ["Store", "Source", "Score", "Type", "ID", "Finding", "Count / sample", "Maps to (posting duty)", "Why it matters", "Fix", "Evidence files (captures/)", "Pillar", "Status"]
    for i, h in enumerate(hdr, 1):
        wf.cell(row=1, column=i, value=h)
    style_header(wf, 1, len(hdr)); wf.row_dimensions[1].height = 30
    KEY = [("F", ["robots", "sitemap", "canonical", "title", "meta", "index", "schema", "structured", "url", "https", "render", "redirect"]),
           ("C", ["description", "image", "alt", "price", "stock", "availab", "variant", "spec", "breadcrumb", "filter", "sort", "search", "catalog", "sku", "product page", "pdp", "taxonomy", "categor", "attribute"]),
           ("B", ["cart", "checkout", "shipping", "return", "review", "rating", "guarantee", "badge", "cta", "merchandis", "best seller", "buy", "order", "quote", "become a customer", "trust", "promo", "collection"]),
           ("T", ["about", "contact", "phone", "address", "privacy", "terms", "social", "accessib", "skip", "footer"]),
           ("N", ["blog", "post", "faq", "guide", "recipe", "locator", "location", "store hours", "news"])]

    def tag(text):
        t = text.lower(); best, bn = "", 0
        for p, kws in KEY:
            n = sum(t.count(k) for k in kws)
            if n > bn: best, bn = p, n
        return best
    fr = 2
    for st in ALL:
        src = "Benchmark" if st["source"] == "index" else "Target"
        for typ, items in (("Strength", st["strengths"]), ("Fix", st["fixes"])):
            for n, txt in enumerate(items, 1):
                vals = [st["company"], src, f'=INDEX(Index!$W$2:$W${LAST},MATCH(A{fr},Index!$C$2:$C${LAST},0))', typ, f"{typ[0]}{n}", txt, "", "", "", "", "", tag(txt), "Published" if src == "Benchmark" else "In review"]
                for col, v in enumerate(vals, 1):
                    wf.cell(row=fr, column=col, value=v)
                finish(wf, fr, len(hdr), wrap_cols=(6, 7, 8, 9, 10, 11), center_cols=(3, 5, 12))
                wf.cell(row=fr, column=3).font = f_green; wf.cell(row=fr, column=3).number_format = "0.0"
                fr += 1
        rv = st.get("review")
        if rv:
            for f in rv["findings"]:
                vals = [st["company"], src, f'=INDEX(Index!$W$2:$W${LAST},MATCH(A{fr},Index!$C$2:$C${LAST},0))', "Finding", f["id"], f["title"], f["count"], f["duty"], f["why"], f["fix"], ", ".join(f["evidence"]), f["pillar"], "Prototyped" if f.get("prototype") else "Documented"]
                for col, v in enumerate(vals, 1):
                    wf.cell(row=fr, column=col, value=v)
                finish(wf, fr, len(hdr), wrap_cols=(6, 7, 8, 9, 10, 11), center_cols=(3, 5, 12))
                wf.cell(row=fr, column=3).font = f_green; wf.cell(row=fr, column=3).number_format = "0.0"
                for col in range(1, len(hdr) + 1):
                    wf.cell(row=fr, column=col).fill = fill_target
                wf.row_dimensions[fr].height = 90
                fr += 1
    wf.freeze_panes = "A2"; wf.auto_filter.ref = f"A1:M{fr - 1}"
    wf.conditional_formatting.add(f"D2:D{fr - 1}", CellIsRule(operator="equal", formula=['"Fix"'], font=Font(name=ARIAL, color=ACCENT, bold=True)))
    wf.conditional_formatting.add(f"D2:D{fr - 1}", CellIsRule(operator="equal", formula=['"Finding"'], font=Font(name=ARIAL, color=COBALT, bold=True)))
    dvs = DataValidation(type="list", formula1='"Documented,Prototyped,Sent,Acknowledged,Shipped by them,Published"', allow_blank=True)
    wf.add_data_validation(dvs); dvs.add(f"M2:M{fr + 300}")

    # ---------------- Impact ledger ----------------
    wl = wb.create_sheet("Impact ledger")
    set_widths(wl, [5, 30, 34, 16, 16, 30, 13, 44, 60])
    hdr = ["#", "Project / employer", "What I did", "Before", "After / result", "Source of the number", "Confidence", "What would make it measured", "Résumé line it supports"]
    for i, h in enumerate(hdr, 1):
        wl.cell(row=1, column=i, value=h)
    style_header(wl, 1, len(hdr)); wl.row_dimensions[1].height = 30
    ledger = [
        ("Universal Wholesale & United Textile (2019–2024)", "Built and ran two storefronts in parallel (WooCommerce + Shopify); catalog maintained by bulk CSV", "—", "14,000+ SKUs maintained: pricing, collections, categories, descriptions, images, SEO metadata", "Résumé; portfolio ‘Counted, not claimed’ tile (Universal Wholesale)", "Stated", "Export the product CSV row count from the store admin and date it", "Owned product setup and catalog data for 14,000+ SKUs"),
        ("United Textile (Shopify)", "Moved the catalog from an offline inventory system into a new Shopify store", "0 products online", "1,000 products live in Shopify", "Portfolio tile (United Textile)", "Stated", "Shopify products export with the launch date", "Migrated 1,000 products from an offline inventory system into Shopify"),
        ("Universal Wholesale / United Textile", "Technical and on-page SEO after the platform migration; redirects and URL structure kept intact", "Not ranking for target queries", "Google pages 1–2 for target queries within the first year", "Résumé (‘improved organic rankings to Google pages 1–2 within the first year’)", "Approximate", "Search Console query report screenshots before/after (if property access remains) — name the queries", "Took organic rankings to Google pages 1–2 within the first year"),
        ("Universal Wholesale", "Led the RepZio B2B platform migration with redirect mapping", "Old platform URLs", "Migration completed with redirects mapped", "Résumé", "Stated", "Count of redirect rules in the map; 404 report after cut-over", "Managed redirects and URL structure through migrations; led the RepZio B2B migration"),
        ("Wild Bill’s Tobacco (2024–2025)", "Administered two corporate sites and the wholesale B2B store; operated the catalog", "—", "~9,000-product catalog operated: accuracy, pricing, descriptions, imagery", "Résumé", "Stated", "WooCommerce product count export, dated", "Operated a ~9,000-product wholesale catalog"),
        ("The Deals Operating System (own build)", "Promotion-accuracy QA workbook + walkthrough for a multi-location retailer (modeled as Northgate Retail Group)", "Promotions checked by hand", "16 checks every promotion passes before it reaches a menu", "Portfolio tile; The Lab", "Measured (the checks exist)", "Count of promotions run through the sheet and the errors it caught in one month", "Built a 16-check promotion QA system"),
        ("ASAS Studio (own product)", "Elementor block library", "—", "3,881 real Elementor blocks (asas.build)", "Portfolio tile — note: the master résumé says ‘4,038 Elementor templates’; reconcile to one number", "Measured", "Count from the ASAS database, dated", "Built ASAS Studio, a library of 3,881 Elementor blocks"),
        ("Independent reviews (Sept 2026)", "Independent, evidence-backed storefront reviews with working prototypes, one per application", "0", "22 reviews published on mousabatarseh.com/reviews (Sept 2026); DiaMedical USA answered one with a formal interview project and an onsite interview", "mousabatarseh.com/reviews; DiaMedical correspondence", "Measured (counts) / Stated (response)", "Log every reply per review on the Outreach sheet; report replies ÷ reviews monthly", "Published 22 independent storefront reviews; one led directly to an interview"),
        ("ShelfMark — Michigan Storefront Benchmark (own build)", "Scored 21 Michigan storefronts on 29 public-page checks", "—", "609 scored observations with evidence; live index + workbook", "ShelfMark workbook README", "Measured", "—", "Built ShelfMark, a 29-check storefront benchmark of 21 Michigan stores"),
        ("Carhartt Reworked review (Sept 2026)", "Read the resale store’s full public feed", "—", "3,545 listings and 124 collections read; 10 findings, 4 working fixes", "mousabatarseh.com/reworked-review", "Measured", "—", "Reviewed Carhartt Reworked end to end: 3,545 listings, 124 collections"),
        ("Blueroot Health review (Sept 28, 2026)", "Four Shopify stores read with the engine", "—", "397 products across four stores; six findings; Fairhaven scored 95/100 (would rank 1st of 21)", "targets/blueroot", "Measured", "—", "—"),
        ("Royal Apparel review (Sept 28, 2026)", "Custom-platform B2B store read with the engine", "—", "1,063 sitemap URLs classified; 0 of 24 pages with a canonical; six findings, five fixes", "targets/royal-apparel", "Measured", "—", "—"),
        ("Federal Fluid Power review (Sept 28, 2026)", "52,000-SKU Magento catalog read with the engine", "—", "54,496 sitemap URLs read; 20,242 keyless product URLs found; six findings, six fixes", "targets/federal-fluid-power", "Measured", "—", "—"),
        ("PageSpeed Insights (a site he runs)", "Performance work on a production site", "—", "Near-perfect PSI scores (Performance, Accessibility, Best Practices, SEO)", "Own statement", "Approximate — site and date not recorded here", "Name the site; save the dated PSI report PDF next to this sheet", "Near-perfect Core Web Vitals on a production site"),
    ]
    for i, row in enumerate(ledger, start=2):
        wl.cell(row=i, column=1, value=i - 1)
        for col, v in enumerate(row, 2):
            wl.cell(row=i, column=col, value=v)
        finish(wl, i, len(hdr), wrap_cols=(2, 3, 4, 5, 6, 8, 9), center_cols=(1, 7))
        wl.row_dimensions[i].height = 60
    dvc = DataValidation(type="list", formula1='"Measured,Stated,Approximate,Measured (counts) / Stated (response),Measured (the checks exist),Approximate — site and date not recorded here"', allow_blank=True)
    wl.add_data_validation(dvc); dvc.add(f"G2:G{len(ledger) + 40}")
    wl.freeze_panes = "C2"
    wl.cell(row=len(ledger) + 3, column=2, value="Measured = a count or export exists. Stated = said on the résumé without an artifact. Approximate = a real outcome without a baseline. Only Measured rows go on the résumé as numbers; Stated rows are written as scope, not results.").font = f_sub
    wl.merge_cells(start_row=len(ledger) + 3, start_column=2, end_row=len(ledger) + 3, end_column=9); wl.cell(row=len(ledger) + 3, column=2).alignment = wrap; wl.row_dimensions[len(ledger) + 3].height = 30

    # ---------------- Outreach ----------------
    wo = wb.create_sheet("Outreach")
    set_widths(wo, [24, 30, 9, 48, 22, 36, 24, 14, 12, 12, 12, 40])
    hdr = ["Company", "Role", "Score", "The one finding in the message", "Contact", "Title", "Channel", "Status", "Sent", "Follow-up", "Reply", "Notes"]
    for i, h in enumerate(hdr, 1):
        wo.cell(row=1, column=i, value=h)
    style_header(wo, 1, len(hdr)); wo.row_dimensions[1].height = 30
    k = 2
    for t in targets:
        rv = t.get("review")
        if not rv:
            continue
        for p in rv["people"]:
            wo.cell(row=k, column=1, value=t["company"]).font = f_bold
            wo.cell(row=k, column=2, value=rv["role"]["title"])
            c = wo.cell(row=k, column=3, value=f'=INDEX(Index!$W$2:$W${LAST},MATCH(A{k},Index!$C$2:$C${LAST},0))'); c.font = f_green; c.number_format = "0.0"
            wo.cell(row=k, column=4, value=rv["findings"][0]["title"])
            wo.cell(row=k, column=5, value=p["name"]); wo.cell(row=k, column=6, value=p["title"]); wo.cell(row=k, column=7, value=p["channel"])
            wo.cell(row=k, column=8, value="Drafted")
            for col in (5, 6, 7, 8, 9, 10, 11, 12):
                wo.cell(row=k, column=col).font = f_input; wo.cell(row=k, column=col).fill = fill_input
            finish(wo, k, len(hdr), wrap_cols=(4, 6, 12), center_cols=(3,))
            wo.row_dimensions[k].height = 32
            k += 1
    dv3 = DataValidation(type="list", formula1='"Not started,Drafted,Sent,Replied,Call booked,Closed"', allow_blank=True)
    wo.add_data_validation(dv3); dv3.add(f"H2:H{k + 200}")
    dv4 = DataValidation(type="list", formula1='"LinkedIn InMail,LinkedIn connection note,Email,Email + LinkedIn connection note,Phone,Indeed,Referral"', allow_blank=True)
    wo.add_data_validation(dv4); dv4.add(f"G2:G{k + 200}")
    wo.freeze_panes = "B2"; wo.auto_filter.ref = f"A1:L{k - 1}"
    wo.cell(row=k + 1, column=1, value="One message per person, within 48 hours of applying: the private review link and ONE finding. Follow up once after 7 days, then move on. Log every reply.").font = f_sub

    # ---------------- Messages ----------------
    wm = wb.create_sheet("Messages")
    wm.sheet_view.showGridLines = False
    set_widths(wm, [3, 110])
    wm["B2"] = "The notes that go with a review"; wm["B2"].font = f_title
    msgs = [
        ("InMail to the hiring manager (≤120 words)", [
            "Subject: [Company]'s storefront — one thing I found, and a full review",
            "Hi [First name] — I applied for the [Role] and, before I did, I read [store] the way the posting describes the job: [scope in a phrase, e.g. four stores, 397 products, no logins].",
            "One finding: [the count — e.g. 3 of the 8 footer Shop links open an empty collection]. The full review, with working fixes and the QA sheet, is here (private link): mousabatarseh.com/[path]/",
            "Everything comes from public pages and it opens with what the store does well. If it's useful, I'd welcome 20 minutes to walk through it. — Mousa Batarseh, Warren MI · hireme@mousabatarseh.com"]),
        ("Connection note (≤300 characters)", [
            "Hi [First name], I applied for [Company]'s [Role] and reviewed [store] first — [one count]. The review with working fixes: mousabatarseh.com/[path]/. Glad to connect. Mousa"]),
        ("Email to a named person (when no LinkedIn route)", [
            "Subject: [Role] application — a review of [store] to go with it",
            "Hi [First name],",
            "I applied for the [Role] on [date]. To show how I'd work rather than say it, I reviewed [store] from its public pages first: [scope]. It opens with three things the store does well, then [N] findings with counts, and [N] working fixes. Private link: mousabatarseh.com/[path]/ — résumé and cover letter are attached.",
            "The one I'd start with: [finding + count + why in one sentence].",
            "Thank you — Mousa Batarseh · (248) 810-1816 · hireme@mousabatarseh.com"]),
        ("Rules", [
            "Name one strength before the finding. Never send a finding that isn't on the Findings sheet. One follow-up after 7 days. Keep the page noindex; the link is private to the person you send it to. Log the row on Outreach the same day."]),
    ]
    r = 4
    for title, lines in msgs:
        wm.cell(row=r, column=2, value=title).font = f_bold; r += 1
        for line in lines:
            c = wm.cell(row=r, column=2, value=line); c.font = f_base; c.alignment = wrap
            wm.row_dimensions[r].height = 15 * (len(line) // 105 + 1); r += 1
        r += 1

    out = os.path.join(HERE, "Review-Engine.xlsx")
    wb.save(out)
    print("saved", out, "| stores", N, "| check rows", CHK_LAST - 1, "| findings rows", fr - 2)


def build_target_workbook(t):
    """The target's own QA workbook: scorecard with live formulas, findings with status, weekly QA sheet, plan."""
    rv = t["review"]
    wb = Workbook()
    ws = wb.active; ws.title = "README"
    rk, n = rank_against_index(t["score"], index)
    readme_block(ws, f"{rv['company']} — review workbook", f"{rv['store_scored']} · {rv['role']['title']} · read {rv['audited_label']} · Mousa Batarseh", [
        ("What this is", f"The spreadsheet behind mousabatarseh.com/{rv['path']}/. Scorecard: the 29 public-page checks with the evidence behind each score (type a new score in the yellow cells and everything recalculates). Findings: the {len(rv['findings'])} role-specific findings with counts, fixes and the files that prove them. Weekly QA: the checks I would run every Monday for this store, pre-filled with today's result. Plan: 30/60/90 days."),
        ("Score", f"{t['score']}/100, grade {t['grade']} ({BANDS[t['grade']]}). Against the ShelfMark Michigan index of {n} stores it would rank {rk} of {n}. Coverage {t['coverage']:.0%} ({t['verified']} of 29 checks verified)."),
        ("Method", "Public pages only: homepage, robots.txt, sitemap, category and product pages, policy pages. Nothing was logged into, carted or submitted. Scores describe a snapshot on the audit date; every finding names the page and the count. No health, legal or financial claim language was evaluated."),
        ("Author", "Mousa Batarseh · hireme@mousabatarseh.com · mousabatarseh.com · linkedin.com/in/mousabatarseh"),
    ])
    # Scorecard
    wsS = wb.create_sheet("Scorecard")
    wsS.sheet_view.showGridLines = False
    set_widths(wsS, [8, 26, 8, 10, 70, 40, 40, 40])
    wsS["A1"] = f"Scorecard — {rv['store_scored']}"; wsS["A1"].font = f_title
    wsS["A2"] = "Type 0, 1, 2 or U in column C to re-score. Pillar %, score, grade and rank recalculate."; wsS["A2"].font = f_sub
    hdr = ["ID", "Check", "Score", "Pillar", f"Evidence ({rv['audited_label']})", "Scores 2 when…", "Scores 1 when…", "Scores 0 when…"]
    for i, h in enumerate(hdr, 1):
        wsS.cell(row=4, column=i, value=h)
    style_header(wsS, 4, 8)
    for i, (cid, name, s2, s1, s0) in enumerate(CHECKS, start=5):
        chk = t["checks"][cid]
        for col, v in enumerate([cid, name, chk["s"], cid[0], chk["e"], s2, s1, s0], 1):
            wsS.cell(row=i, column=col, value=v)
        finish(wsS, i, 8, wrap_cols=(5, 6, 7, 8), center_cols=(1, 3, 4))
        sc = wsS.cell(row=i, column=3); sc.font = f_input; sc.fill = fill_input
        wsS.row_dimensions[i].height = 44
    S1, S2 = 5, 5 + len(CHECKS) - 1
    dv = DataValidation(type="list", formula1='"0,1,2,U"', allow_blank=False); wsS.add_data_validation(dv); dv.add(f"C{S1}:C{S2}")
    R0 = S2 + 2
    wsS.cell(row=R0, column=1, value="Results").font = f_title
    for i, h in enumerate(["Pillar", "Name", "Weight", "Verified", "Points", "Pillar %"], 1):
        wsS.cell(row=R0 + 1, column=i, value=h)
    style_header(wsS, R0 + 1, 6)
    SR = f"$C${S1}:$C${S2}"; SP = f"$D${S1}:$D${S2}"
    for k, (pid, name, w, _) in enumerate(PILLARS):
        rr = R0 + 2 + k
        wsS.cell(row=rr, column=1, value=pid); wsS.cell(row=rr, column=2, value=name); wsS.cell(row=rr, column=3, value=w)
        wsS.cell(row=rr, column=4, value=f'=COUNTIFS({SP},A{rr},{SR},">=0")')
        wsS.cell(row=rr, column=5, value=f'=SUMIFS({SR},{SP},A{rr})')
        c = wsS.cell(row=rr, column=6, value=f'=IF(D{rr}=0,"",E{rr}/(2*D{rr}))'); c.number_format = "0%"
        finish(wsS, rr, 6, center_cols=(1, 3, 4, 5, 6))
    P1 = R0 + 2; RS = P1 + 7
    den = "+".join(f"IF(D{P1+k}>0,C{P1+k},0)" for k in range(5)); num = "+".join(f"IF(D{P1+k}>0,C{P1+k}*F{P1+k},0)" for k in range(5))
    bench = ",".join(str(r["score"]) for r in index)
    rows_ = [("Score", f'=IF(({den})=0,"",ROUND(100*({num})/({den}),1))', "0.0"),
             ("Grade", f'=IF(B{RS}="","",IF(B{RS}>=85,"A",IF(B{RS}>=70,"B",IF(B{RS}>=55,"C",IF(B{RS}>=40,"D","F")))))', None),
             ("Verified checks", f"=SUM(D{P1}:D{P1+4})", "0"),
             ("Coverage", f"=B{RS+2}/29", "0%"),
             (f"Would rank on the Michigan index (of {n})", f'=IF(B{RS}="","",1+SUMPRODUCT(--({{{bench}}}>B{RS})))', "0"),
             ("Michigan index average", round(sum(r["score"] for r in index) / n, 1), "0.0"),
             ("Weakest pillar", f'=INDEX(B{P1}:B{P1+4},MATCH(MIN(F{P1}:F{P1+4}),F{P1}:F{P1+4},0))', None)]
    for k, (lab, fml, fmt) in enumerate(rows_):
        rr = RS + k
        wsS.cell(row=rr, column=1, value=lab).font = f_bold
        c = wsS.cell(row=rr, column=2, value=fml); c.font = Font(name=ARIAL, size=14 if k == 0 else 10, bold=k < 2, color=ACCENT if k == 0 else INK)
        if fmt: c.number_format = fmt
        wsS.merge_cells(start_row=rr, start_column=2, end_row=rr, end_column=3)
        for col in (1, 2, 3): wsS.cell(row=rr, column=col).border = box
    wsS.freeze_panes = "A5"
    # Findings
    wf = wb.create_sheet("Findings")
    set_widths(wf, [6, 46, 44, 34, 56, 64, 50, 9, 14])
    hdr = ["ID", "Finding", "Count / sample", "Maps to (posting duty)", "Why it matters", "Fix", "Evidence (captures/)", "Pillar", "Status"]
    for i, h in enumerate(hdr, 1):
        wf.cell(row=1, column=i, value=h)
    style_header(wf, 1, len(hdr)); wf.row_dimensions[1].height = 30
    r = 2
    for s in rv["strengths"]:
        for col, v in enumerate(["S", s["title"], s["text"], "", "", "", "", "", "Strength"], 1):
            wf.cell(row=r, column=col, value=v)
        finish(wf, r, len(hdr), wrap_cols=(2, 3), center_cols=(1, 8)); wf.row_dimensions[r].height = 48; r += 1
    for f in rv["findings"]:
        for col, v in enumerate([f["id"], f["title"], f["count"], f["duty"], f["why"], f["fix"], ", ".join(f["evidence"]), f["pillar"], "Prototyped" if f.get("prototype") else "Documented"], 1):
            wf.cell(row=r, column=col, value=v)
        finish(wf, r, len(hdr), wrap_cols=(2, 3, 4, 5, 6, 7), center_cols=(1, 8)); wf.row_dimensions[r].height = 120
        c = wf.cell(row=r, column=9); c.font = f_input; c.fill = fill_input
        r += 1
    for a in rv.get("also", []):
        for col, v in enumerate(["·", a, "", "", "", "", "", "", "Noted"], 1):
            wf.cell(row=r, column=col, value=v)
        finish(wf, r, len(hdr), wrap_cols=(2,), center_cols=(1,)); wf.row_dimensions[r].height = 30; r += 1
    dvs = DataValidation(type="list", formula1='"Strength,Documented,Prototyped,Noted,Fixed,Verified fixed"', allow_blank=True)
    wf.add_data_validation(dvs); dvs.add(f"I2:I{r + 100}")
    wf.freeze_panes = "B2"; wf.auto_filter.ref = f"A1:I{r - 1}"
    # Weekly QA
    wq = wb.create_sheet("Weekly QA")
    set_widths(wq, [5, 44, 40, 22, 14, 30, 30])
    wq["A1"] = "Weekly QA — the Monday checks for this store"; wq["A1"].font = f_title
    wq["A2"] = "One row per check. Type the result each week; keep the history as new columns."; wq["A2"].font = f_sub
    hdr = ["#", "Check", "How (public page / feed)", "Owner", "Today", "Result today", "Threshold / action"]
    for i, h in enumerate(hdr, 1):
        wq.cell(row=4, column=i, value=h)
    style_header(wq, 4, len(hdr))
    qa = rv.get("weekly_qa") or []
    if not qa:
        qa = [[f["duty"], f"Re-check: {f['title']}", "Ops", f["count"], "Fix per Findings"] for f in rv["findings"]]
    for i, row in enumerate(qa, start=5):
        wq.cell(row=i, column=1, value=i - 4)
        wq.cell(row=i, column=2, value=row[0]); wq.cell(row=i, column=3, value=row[1]); wq.cell(row=i, column=4, value=row[2])
        wq.cell(row=i, column=5, value=TODAY); wq.cell(row=i, column=6, value=row[3]); wq.cell(row=i, column=7, value=row[4])
        finish(wq, i, len(hdr), wrap_cols=(2, 3, 6, 7), center_cols=(1, 5))
        for col in (4, 6): wq.cell(row=i, column=col).font = f_input; wq.cell(row=i, column=col).fill = fill_input
        wq.row_dimensions[i].height = 44
    wq.freeze_panes = "C5"
    # Plan
    wp = wb.create_sheet("Plan")
    set_widths(wp, [12, 100])
    wp["A1"] = "30 / 60 / 90 days"; wp["A1"].font = f_title
    r = 3
    for k in ("30", "60", "90"):
        wp.cell(row=r, column=1, value=f"Day {k}").font = f_bold
        for item in rv["plan"][k]:
            c = wp.cell(row=r, column=2, value=item); c.font = f_base; c.alignment = wrap; wp.row_dimensions[r].height = 30; r += 1
        r += 1
    outdir = os.path.join(HERE, "targets", t["slug"], "site", "files")
    os.makedirs(outdir, exist_ok=True)
    out = os.path.join(outdir, os.path.basename(rv["files"]["workbook"]))
    wb.save(out)
    print("saved", out)


if __name__ == "__main__":
    build_master()
    for t in targets:
        if t.get("review"):
            build_target_workbook(t)
