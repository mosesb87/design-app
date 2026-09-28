import json
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import CellIsRule, DataBarRule, FormulaRule
from openpyxl.comments import Comment
from model import build_dataset, PILLARS, CHECKS, CHECK_IDS, WEIGHTS, GRADES, BANDS

rows = build_dataset()
N = len(rows)
LAST = N + 1  # last data row on Index (header is row 1)

# ---------- styles ----------
ARIAL = "Arial"
INK = "1F1D1A"; PAPER = "F4F1EA"; ACCENT = "C8451F"; MUTED = "6B6660"; LINE = "D9D4CA"
f_base = Font(name=ARIAL, size=10, color=INK)
f_bold = Font(name=ARIAL, size=10, bold=True, color=INK)
f_head = Font(name=ARIAL, size=10, bold=True, color="FFFFFF")
f_title = Font(name=ARIAL, size=18, bold=True, color=INK)
f_sub = Font(name=ARIAL, size=11, color=MUTED)
f_input = Font(name=ARIAL, size=10, color="0000FF")
f_link = Font(name=ARIAL, size=10, color="0000FF", underline="single")
f_green = Font(name=ARIAL, size=10, color="008000")
fill_head = PatternFill("solid", fgColor=INK)
fill_sub = PatternFill("solid", fgColor="E9E4D9")
fill_input = PatternFill("solid", fgColor="FFFF00")
fill_soft = PatternFill("solid", fgColor=PAPER)
thin = Side(style="thin", color=LINE)
box = Border(left=thin, right=thin, top=thin, bottom=thin)
wrap = Alignment(wrap_text=True, vertical="top")
center = Alignment(horizontal="center", vertical="center")

def style_header(ws, row, ncols, fill=fill_head, font=f_head):
    for c in range(1, ncols + 1):
        cell = ws.cell(row=row, column=c)
        cell.fill = fill; cell.font = font; cell.border = box
        cell.alignment = Alignment(wrap_text=True, vertical="center")

def set_widths(ws, widths):
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w

def body(ws, r1, r2, c1, c2, font=f_base, align=wrap):
    for r in range(r1, r2 + 1):
        for c in range(c1, c2 + 1):
            cell = ws.cell(row=r, column=c)
            if cell.font == Font():  # untouched default
                cell.font = font
            cell.border = box
            if align is not None and cell.alignment.wrap_text is None:
                cell.alignment = align

wb = Workbook()

# =====================================================================
# README
# =====================================================================
ws = wb.active; ws.title = "README"
ws.sheet_view.showGridLines = False
set_widths(ws, [3, 26, 90])
ws["B2"] = "ShelfMark — Michigan Storefront Benchmark"; ws["B2"].font = f_title
ws["B3"] = "Scoring model, audit data and outreach tracker · Vol. 1 · September 2026 · by Moses Batarseh"; ws["B3"].font = f_sub
lines = [
    ("What this is", "A repeatable scoring system for e-commerce storefronts: 29 public-page checks across 5 pillars, weighted into one 0–100 ShelfMark Score. Volume 1 audits 21 Michigan companies (retail chains, wholesale/B2B distributors, food & drink makers, brands & makers) from their public pages only, on 2026-09-28."),
    ("How a score is built", "Each check is scored 2 (meets), 1 (partly), 0 (fails) or U (unverified — could not be observed from public pages). Pillar % = verified points ÷ (2 × verified checks). ShelfMark Score = weighted average of pillar % (weights on the Rubric sheet: Findability 25, Catalog 30, Buy 20, Trust 15, Content 10). A pillar with no verified checks drops out and the other weights renormalise. U never counts for or against a store; Coverage (verified ÷ 29) and Confidence show how much was observable."),
    ("Grades", "A ≥ 85 Benchmark · B ≥ 70 Strong · C ≥ 55 Developing · D ≥ 40 Needs work · F < 40 At risk. Thresholds are editable on the Rubric sheet."),
    ("Sheets", "Index — the ranked table (all formulas; nothing hardcoded). Checks — 609 scored observations with evidence, the only place scores are typed. Rubric — pillars, weights and the 2/1/0 definitions. Score a Store — type 29 scores for any store and see its score, grade and where it would rank. Findings — 3 strengths + 3 fixes per store. Insights — averages by segment and platform, pass rates per check. Outreach — one row per company to turn report cards into conversations. Outreach message — the note that goes with a report card."),
    ("How to add a store", "1) Append 29 rows to Checks (Store, Pillar letter, Check ID, Check name, Score, Evidence). 2) Add one row to Index with the store name in column B and copy the formulas of the row above across. 3) Extend the $2:$22 ranges in Index (Rank), Score a Store and Insights by one row. Ranks, grades and averages update automatically."),
    ("Legend", "Blue text on a yellow fill = cells meant to be typed in (scores, contacts, statuses). Black = formulas. Green = a value pulled from another sheet."),
    ("Method limits", "Audits use only publicly fetchable pages (homepage, robots.txt, sitemap, one category, two product pages, contact/about). Client-rendered pages, bot challenges and checkout flows are marked U rather than guessed. Structured data (F7) is U unless schema was visible in served HTML. Scores describe a snapshot on the audit date."),
    ("Source", "Audit evidence: each Checks row quotes the page or text observed, with the sampled URLs listed on the Index sheet. Company HQ cities were confirmed from each company's own site."),
]
r = 5
for k, v in lines:
    ws.cell(row=r, column=2, value=k).font = f_bold
    ws.cell(row=r, column=2).alignment = wrap
    c = ws.cell(row=r, column=3, value=v); c.font = f_base; c.alignment = wrap
    ws.row_dimensions[r].height = max(30, 15 * (len(v) // 95 + 1))
    r += 1
ws.cell(row=r + 1, column=2, value="Author").font = f_bold
ws.cell(row=r + 1, column=3, value="Moses Batarseh · hireme@mousabatarseh.com · mousabatarseh.com · linkedin.com/in/mousabatarseh").font = f_base

# =====================================================================
# Rubric (weights + thresholds are inputs referenced by formulas)
# =====================================================================
wr = wb.create_sheet("Rubric")
wr.sheet_view.showGridLines = False
set_widths(wr, [8, 24, 10, 60, 48, 48, 48])
wr["A1"] = "Pillars & weights"; wr["A1"].font = f_title
hdr = ["Pillar", "Name", "Weight", "Question the pillar answers"]
for i, h in enumerate(hdr, 1):
    wr.cell(row=2, column=i, value=h)
style_header(wr, 2, 4)
for i, (pid, name, w, q) in enumerate(PILLARS, start=3):
    wr.cell(row=i, column=1, value=pid).font = f_base
    wr.cell(row=i, column=2, value=name).font = f_base
    c = wr.cell(row=i, column=3, value=w); c.font = f_input; c.fill = fill_input
    wr.cell(row=i, column=4, value=q).font = f_base
    for col in range(1, 5): wr.cell(row=i, column=col).border = box
wr.cell(row=8, column=2, value="Total").font = f_bold
wr.cell(row=8, column=3, value="=SUM(C3:C7)").font = f_bold
wr.cell(row=8, column=4, value="Weights are relative; they need not sum to 100.").font = f_sub
# Grade thresholds
wr["A10"] = "Grade thresholds (score ≥)"; wr["A10"].font = f_bold
for i, (cut, g) in enumerate(GRADES[:4], start=11):
    wr.cell(row=i, column=1, value=g).font = f_base
    wr.cell(row=i, column=2, value=BANDS[g]).font = f_base
    c = wr.cell(row=i, column=3, value=cut); c.font = f_input; c.fill = fill_input
    for col in range(1, 4): wr.cell(row=i, column=col).border = box
wr.cell(row=15, column=1, value="F").font = f_base; wr.cell(row=15, column=2, value=BANDS["F"]).font = f_base
wr.cell(row=15, column=3, value="below D").font = f_sub
for col in range(1, 4): wr.cell(row=15, column=col).border = box
wr["A16"] = "Confidence: High = 26+ verified checks, Medium = 20–25, Low = under 20 (of 29)."; wr["A16"].font = f_sub
# Checks table
wr["A18"] = "The 29 checks"; wr["A18"].font = f_title
hdr = ["ID", "Check", "Pillar", "Scores 2 when…", "Scores 1 when…", "Scores 0 when…"]
for i, h in enumerate(hdr, 1):
    wr.cell(row=19, column=i, value=h)
style_header(wr, 19, 6)
for i, (cid, name, s2, s1, s0) in enumerate(CHECKS, start=20):
    vals = [cid, name, cid[0], s2, s1, s0]
    for col, v in enumerate(vals, 1):
        c = wr.cell(row=i, column=col, value=v); c.font = f_base; c.border = box; c.alignment = wrap
    wr.row_dimensions[i].height = 42
wr.freeze_panes = "A20"
RUB_LAST = 19 + len(CHECKS)

# =====================================================================
# Checks (long format)
# =====================================================================
wc = wb.create_sheet("Checks")
set_widths(wc, [28, 8, 9, 24, 8, 110])
hdr = ["Store", "Pillar", "Check ID", "Check", "Score", "Evidence (what was observed, 2026-09-28)"]
for i, h in enumerate(hdr, 1):
    wc.cell(row=1, column=i, value=h)
style_header(wc, 1, 6)
cr = 2
for st in rows:
    for cid, name, *_ in CHECKS:
        chk = st["checks"][cid]
        vals = [st["company"], cid[0], cid, name, chk["s"], chk["e"]]
        for col, v in enumerate(vals, 1):
            c = wc.cell(row=cr, column=col, value=v); c.font = f_base; c.border = box
            c.alignment = Alignment(vertical="top", wrap_text=(col == 6))
        sc = wc.cell(row=cr, column=5); sc.font = f_input; sc.fill = fill_input; sc.alignment = center
        cr += 1
CHK_LAST = cr - 1
dv = DataValidation(type="list", formula1='"0,1,2,U"', allow_blank=False, showErrorMessage=True,
                    errorTitle="Score", error="Enter 0, 1, 2 or U")
wc.add_data_validation(dv); dv.add(f"E2:E{CHK_LAST + 300}")
wc.freeze_panes = "A2"; wc.auto_filter.ref = f"A1:F{CHK_LAST}"

# =====================================================================
# Index
# =====================================================================
wi = wb.create_sheet("Index", 1)
wi.sheet_view.showGridLines = False
cols = ["Rank", "Store", "Segment", "Platform", "HQ", "Sells", "URL", "Audited",
        "Findability %", "F n", "Catalog %", "C n", "Buy %", "B n", "Trust %", "T n", "Content %", "N n",
        "Verified (of 29)", "Coverage", "Confidence", "ShelfMark Score", "Grade", "Band",
        "Category sampled", "Product sampled 1", "Product sampled 2", "Notes"]
for i, h in enumerate(cols, 1):
    wi.cell(row=1, column=i, value=h)
style_header(wi, 1, len(cols))
wi.row_dimensions[1].height = 32
set_widths(wi, [6, 28, 18, 19, 20, 30, 34, 11, 12, 5, 11, 5, 9, 5, 9, 5, 11, 5, 10, 10, 11, 12, 7, 12, 40, 40, 40, 50])
W = {p: f"Rubric!$C${3 + i}" for i, (p, *_) in enumerate(PILLARS)}  # weight cell refs
PC = {"F": ("I", "J"), "C": ("K", "L"), "B": ("M", "N"), "T": ("O", "P"), "N": ("Q", "R")}
CHK = f"Checks!$A$2:$A${CHK_LAST + 300}"; CHP = f"Checks!$B$2:$B${CHK_LAST + 300}"; CHS = f"Checks!$E$2:$E${CHK_LAST + 300}"
for idx, st in enumerate(rows):
    r = idx + 2
    wi.cell(row=r, column=1, value=f'=IF(V{r}="","",RANK(V{r},$V$2:$V${LAST},0))')
    wi.cell(row=r, column=2, value=st["company"]).font = f_bold
    wi.cell(row=r, column=3, value=st["segment"])
    wi.cell(row=r, column=4, value=st["platform"])
    wi.cell(row=r, column=5, value=st["hq_city"])
    wi.cell(row=r, column=6, value=st.get("sells", ""))
    u = wi.cell(row=r, column=7, value=st["url"]); u.hyperlink = st["url"]; u.font = f_link
    wi.cell(row=r, column=8, value=st["audited_at"])
    for p, (pcol, ncol) in PC.items():
        wi[f"{pcol}{r}"] = (f'=IFERROR(SUMIFS({CHS},{CHK},$B{r},{CHP},"{p}")'
                            f'/(2*COUNTIFS({CHK},$B{r},{CHP},"{p}",{CHS},">=0")),"")')
        wi[f"{pcol}{r}"].number_format = "0%"
        wi[f"{ncol}{r}"] = f'=COUNTIFS({CHK},$B{r},{CHP},"{p}",{CHS},">=0")'
        wi[f"{ncol}{r}"].font = Font(name=ARIAL, size=8, color=MUTED)
    wi[f"S{r}"] = f"=J{r}+L{r}+N{r}+P{r}+R{r}"
    wi[f"T{r}"] = f"=S{r}/29"; wi[f"T{r}"].number_format = "0%"
    wi[f"U{r}"] = f'=IF(S{r}>=26,"High",IF(S{r}>=20,"Medium","Low"))'
    num = "+".join(f'IF({PC[p][1]}{r}>0,{W[p]}*{PC[p][0]}{r},0)' for p in PC)
    den = "+".join(f'IF({PC[p][1]}{r}>0,{W[p]},0)' for p in PC)
    wi[f"V{r}"] = f'=IF(({den})=0,"",ROUND(100*({num})/({den}),1))'
    wi[f"V{r}"].font = Font(name=ARIAL, size=11, bold=True, color=INK); wi[f"V{r}"].number_format = "0.0"
    wi[f"W{r}"] = (f'=IF(V{r}="","",IF(V{r}>=Rubric!$C$11,"A",IF(V{r}>=Rubric!$C$12,"B",'
                   f'IF(V{r}>=Rubric!$C$13,"C",IF(V{r}>=Rubric!$C$14,"D","F")))))')
    wi[f"W{r}"].font = f_bold; wi[f"W{r}"].alignment = center
    wi[f"X{r}"] = (f'=IF(W{r}="","",IF(W{r}="A","Benchmark",IF(W{r}="B","Strong",IF(W{r}="C","Developing",'
                   f'IF(W{r}="D","Needs work","At risk")))))')
    su = st.get("sample_urls", {})
    for col, key in zip((25, 26, 27), ("category", "product1", "product2")):
        v = su.get(key) or ""
        c = wi.cell(row=r, column=col, value=v)
        if v: c.hyperlink = v; c.font = f_link
    wi.cell(row=r, column=28, value=st.get("notes", ""))
    for col in range(1, len(cols) + 1):
        c = wi.cell(row=r, column=col); c.border = box
        if c.font == Font(): c.font = f_base
        c.alignment = Alignment(vertical="center", wrap_text=col in (6, 28), horizontal="center" if col in (1, 8, 10, 12, 14, 16, 18, 19, 20, 21, 23) else None)
    wi.row_dimensions[r].height = 30
wi.freeze_panes = "C2"; wi.auto_filter.ref = f"A1:AB{LAST}"
# conditional formatting: score bar, grade colours
wi.conditional_formatting.add(f"V2:V{LAST}", DataBarRule(start_type="num", start_value=0, end_type="num", end_value=100, color=ACCENT))
for g, colr in (("A", "C6EFCE"), ("B", "E2F0D9"), ("C", "FFF2CC"), ("D", "FCE4D6"), ("F", "F8CBAD")):
    wi.conditional_formatting.add(f"W2:W{LAST}", CellIsRule(operator="equal", formula=[f'"{g}"'], fill=PatternFill("solid", fgColor=colr)))
for col in ("I", "K", "M", "O", "Q"):
    wi.conditional_formatting.add(f"{col}2:{col}{LAST}", DataBarRule(start_type="num", start_value=0, end_type="num", end_value=1, color="8A8377"))
wi.cell(row=LAST + 2, column=2, value="Index average").font = f_bold
wi.cell(row=LAST + 2, column=22, value=f"=ROUND(AVERAGE(V2:V{LAST}),1)").font = f_bold
wi.cell(row=LAST + 3, column=2, value="Index median").font = f_bold
wi.cell(row=LAST + 3, column=22, value=f"=MEDIAN(V2:V{LAST})").font = f_bold
wi.cell(row=LAST + 5, column=2, value="How to read: pillar % = verified points ÷ (2 × verified checks); 'n' = verified checks in that pillar (of 8/8/5/5/3). Score = weighted average of pillar % (weights on Rubric). U never counts for or against a store — see Coverage and Confidence.").font = f_sub
wi.merge_cells(start_row=LAST + 5, start_column=2, end_row=LAST + 5, end_column=12)
wi.cell(row=LAST + 5, column=2).alignment = wrap; wi.row_dimensions[LAST + 5].height = 45

# =====================================================================
# Score a Store (calculator)
# =====================================================================
wsS = wb.create_sheet("Score a Store", 2)
wsS.sheet_view.showGridLines = False
set_widths(wsS, [8, 26, 8, 10, 46, 40, 30, 30])
wsS["A1"] = "Score a Store"; wsS["A1"].font = f_title
wsS["A2"] = "Type 0, 1, 2 or U in the yellow cells (column C). Everything else recalculates. The example values below are illustrative — replace them with your own audit."; wsS["A2"].font = f_sub
wsS["A4"] = "Store"; wsS["A4"].font = f_bold
c = wsS["B4"]; c.value = "Example Outfitters (replace me)"; c.font = f_input; c.fill = fill_input
wsS["A5"] = "URL"; wsS["A5"].font = f_bold
c = wsS["B5"]; c.value = "https://example.com"; c.font = f_input; c.fill = fill_input
wsS["A6"] = "Audited"; wsS["A6"].font = f_bold
c = wsS["B6"]; c.value = "2026-09-28"; c.font = f_input; c.fill = fill_input
hdr = ["ID", "Check", "Score", "Pillar", "Evidence / notes", "Scores 2 when…", "Scores 1 when…", "Scores 0 when…"]
for i, h in enumerate(hdr, 1):
    wsS.cell(row=8, column=i, value=h)
style_header(wsS, 8, 8)
example = {"F1": 2, "F2": 1, "F3": 2, "F4": 1, "F5": 0, "F6": 2, "F7": "U", "F8": 2,
           "C1": 2, "C2": 1, "C3": 1, "C4": 2, "C5": 1, "C6": 1, "C7": 2, "C8": 2,
           "B1": 2, "B2": 1, "B3": "U", "B4": 2, "B5": 1,
           "T1": 2, "T2": 1, "T3": 2, "T4": 2, "T5": 1,
           "N1": 0, "N2": 1, "N3": 2}
for i, (cid, name, *_) in enumerate(CHECKS, start=9):
    wsS.cell(row=i, column=1, value=cid); wsS.cell(row=i, column=2, value=name)
    sc = wsS.cell(row=i, column=3, value=example[cid]); sc.font = f_input; sc.fill = fill_input; sc.alignment = center
    wsS.cell(row=i, column=4, value=f'=LEFT(A{i},1)')
    ev = wsS.cell(row=i, column=5, value="Example — replace"); ev.font = f_input; ev.fill = fill_input
    rr = 20 + (i - 9)  # rubric row for this check
    for col, rcol in zip((6, 7, 8), ("D", "E", "F")):
        wsS.cell(row=i, column=col, value=f"=Rubric!{rcol}{rr}").font = f_green
    for col in range(1, 9):
        c = wsS.cell(row=i, column=col); c.border = box
        if c.font == Font(): c.font = f_base
        c.alignment = Alignment(vertical="top", wrap_text=col >= 5, horizontal="center" if col in (1, 3, 4) else None)
    wsS.row_dimensions[i].height = 40
S_FIRST, S_LAST = 9, 9 + len(CHECKS) - 1
dv2 = DataValidation(type="list", formula1='"0,1,2,U"', allow_blank=False)
wsS.add_data_validation(dv2); dv2.add(f"C{S_FIRST}:C{S_LAST}")
# results block
R0 = S_LAST + 2
wsS.cell(row=R0, column=1, value="Results").font = f_title
hdr = ["Pillar", "Name", "Weight", "Verified", "Points", "Pillar %"]
for i, h in enumerate(hdr, 1):
    wsS.cell(row=R0 + 1, column=i, value=h)
style_header(wsS, R0 + 1, 6)
SR = f"$C${S_FIRST}:$C${S_LAST}"; SP = f"$D${S_FIRST}:$D${S_LAST}"
for k, (pid, name, w, _) in enumerate(PILLARS):
    rr = R0 + 2 + k
    wsS.cell(row=rr, column=1, value=pid); wsS.cell(row=rr, column=2, value=name)
    wsS.cell(row=rr, column=3, value=f"=Rubric!C{3 + k}").font = f_green
    wsS.cell(row=rr, column=4, value=f'=COUNTIFS({SP},A{rr},{SR},">=0")')
    wsS.cell(row=rr, column=5, value=f'=SUMIFS({SR},{SP},A{rr})')
    c = wsS.cell(row=rr, column=6, value=f'=IF(D{rr}=0,"",E{rr}/(2*D{rr}))'); c.number_format = "0%"
    for col in range(1, 7):
        cc = wsS.cell(row=rr, column=col); cc.border = box
        if cc.font == Font(): cc.font = f_base
P1, P5 = R0 + 2, R0 + 6
RS = P5 + 2
labels = [
    ("ShelfMark Score", f'=IF(SUMPRODUCT(--(D{P1}:D{P5}>0),C{P1}:C{P5})=0,"",ROUND(100*SUMPRODUCT(--(D{P1}:D{P5}>0),C{P1}:C{P5},IF(D{P1}:D{P5}>0,F{P1}:F{P5},0))/SUMPRODUCT(--(D{P1}:D{P5}>0),C{P1}:C{P5}),1))', "0.0"),
    ("Grade", f'=IF(B{RS}="","",IF(B{RS}>=Rubric!$C$11,"A",IF(B{RS}>=Rubric!$C$12,"B",IF(B{RS}>=Rubric!$C$13,"C",IF(B{RS}>=Rubric!$C$14,"D","F")))))', None),
    ("Verified checks", f'=SUM(D{P1}:D{P5})', "0"),
    ("Coverage", f'=B{RS + 2}/29', "0%"),
    ("Confidence", f'=IF(B{RS + 2}>=26,"High",IF(B{RS + 2}>=20,"Medium","Low"))', None),
    ("Would rank (of the index)", f'=IF(B{RS}="","",COUNTIF(Index!$V$2:$V${LAST},">"&B{RS})+1&" of "&COUNT(Index!$V$2:$V${LAST}))', None),
    ("Index average", f'=ROUND(AVERAGE(Index!$V$2:$V${LAST}),1)', "0.0"),
    ("Gap to index leader", f'=IF(B{RS}="","",ROUND(MAX(Index!$V$2:$V${LAST})-B{RS},1))', "0.0"),
    ("Weakest pillar", f'=IF(COUNT(F{P1}:F{P5})=0,"",INDEX(B{P1}:B{P5},MATCH(MIN(F{P1}:F{P5}),F{P1}:F{P5},0)))', None),
]
for k, (lab, fml, fmt) in enumerate(labels):
    rr = RS + k
    wsS.cell(row=rr, column=1, value=lab).font = f_bold
    c = wsS.cell(row=rr, column=2, value=fml); c.font = Font(name=ARIAL, size=12 if k < 2 else 10, bold=k < 2, color=INK)
    if fmt: c.number_format = fmt
    for col in (1, 2): wsS.cell(row=rr, column=col).border = box
wsS.merge_cells(start_row=RS, start_column=1, end_row=RS, end_column=1)
# the SUMPRODUCT with IF needs array eval in Excel; use a safer form instead:
wsS.cell(row=RS, column=2, value=(
    f'=IF(({"+".join(f"IF(D{P1+k}>0,C{P1+k},0)" for k in range(5))})=0,"",'
    f'ROUND(100*({"+".join(f"IF(D{P1+k}>0,C{P1+k}*F{P1+k},0)" for k in range(5))})'
    f'/({"+".join(f"IF(D{P1+k}>0,C{P1+k},0)" for k in range(5))}),1))'))
wsS.cell(row=RS, column=2).font = Font(name=ARIAL, size=14, bold=True, color=ACCENT)
wsS.freeze_panes = "A9"

# =====================================================================
# Findings
# =====================================================================
wf = wb.create_sheet("Findings")
set_widths(wf, [28, 7, 9, 11, 4, 110, 12])
hdr = ["Store", "Rank", "Score", "Type", "#", "Finding", "Pillar (auto-tag)"]
for i, h in enumerate(hdr, 1):
    wf.cell(row=1, column=i, value=h)
style_header(wf, 1, 7)
KEY = [("F", ["robots", "sitemap", "canonical", "title", "meta", "index", "schema", "structured", "url", "search console", "https", "render"]),
       ("C", ["description", "image", "alt", "price", "stock", "availab", "variant", "spec", "breadcrumb", "filter", "sort", "search", "catalog", "sku", "product page", "pdp", "taxonomy", "categor"]),
       ("B", ["cart", "checkout", "shipping", "return", "review", "rating", "guarantee", "badge", "cta", "merchandis", "best seller", "buy", "order", "quote", "become a customer", "trust"]),
       ("T", ["about", "contact", "phone", "address", "privacy", "terms", "social", "accessib", "skip"]),
       ("N", ["blog", "post", "faq", "guide", "recipe", "locator", "location", "store hours", "news"])]
def tag(text):
    t = text.lower()
    best, bn = "", 0
    for p, kws in KEY:
        n = sum(t.count(k) for k in kws)
        if n > bn: best, bn = p, n
    return best
fr = 2
for st in rows:
    for typ, items in (("Strength", st["strengths"]), ("Fix", st["fixes"])):
        for n, txt in enumerate(items, 1):
            vals = [st["company"], f'=INDEX(Index!$A$2:$A${LAST},MATCH(A{fr},Index!$B$2:$B${LAST},0))',
                    f'=INDEX(Index!$V$2:$V${LAST},MATCH(A{fr},Index!$B$2:$B${LAST},0))', typ, n, txt, tag(txt)]
            for col, v in enumerate(vals, 1):
                c = wf.cell(row=fr, column=col, value=v); c.font = f_green if col in (2, 3) else f_base; c.border = box
                c.alignment = Alignment(vertical="top", wrap_text=(col == 6), horizontal="center" if col in (2, 3, 5, 7) else None)
            wf.cell(row=fr, column=3).number_format = "0.0"
            wf.row_dimensions[fr].height = 30
            fr += 1
wf.freeze_panes = "A2"; wf.auto_filter.ref = f"A1:G{fr - 1}"
wf.conditional_formatting.add(f"D2:D{fr - 1}", CellIsRule(operator="equal", formula=['"Fix"'], font=Font(name=ARIAL, color=ACCENT, bold=True)))

# =====================================================================
# Insights
# =====================================================================
wn = wb.create_sheet("Insights")
wn.sheet_view.showGridLines = False
set_widths(wn, [24, 10, 12, 12, 12, 12, 12, 12, 12])
wn["A1"] = "Insights — computed live from Index and Checks"; wn["A1"].font = f_title
IDX_SEG = f"Index!$C$2:$C${LAST}"; IDX_PLT = f"Index!$D$2:$D${LAST}"; IDX_SC = f"Index!$V$2:$V${LAST}"
wn["A3"] = "By segment"; wn["A3"].font = f_bold
for i, h in enumerate(["Segment", "Stores", "Avg score", "Best", "Lowest"], 1):
    wn.cell(row=4, column=i, value=h)
style_header(wn, 4, 5)
segs = sorted({r["segment"] for r in rows})
for k, s in enumerate(segs, start=5):
    wn.cell(row=k, column=1, value=s)
    wn.cell(row=k, column=2, value=f'=COUNTIF({IDX_SEG},A{k})')
    wn.cell(row=k, column=3, value=f'=IFERROR(ROUND(AVERAGEIF({IDX_SEG},A{k},{IDX_SC}),1),"")')
    wn.cell(row=k, column=4, value=f'=IFERROR(_xlfn.MAXIFS({IDX_SC},{IDX_SEG},A{k}),"")')
    wn.cell(row=k, column=5, value=f'=IFERROR(_xlfn.MINIFS({IDX_SC},{IDX_SEG},A{k}),"")')
    for col in range(1, 6):
        c = wn.cell(row=k, column=col); c.border = box; c.font = f_base
        if col >= 3: c.number_format = "0.0"
r0 = 5 + len(segs) + 2
wn.cell(row=r0, column=1, value="By platform").font = f_bold
for i, h in enumerate(["Platform", "Stores", "Avg score", "Best", "Lowest"], 1):
    wn.cell(row=r0 + 1, column=i, value=h)
style_header(wn, r0 + 1, 5)
plats = sorted({r["platform"] for r in rows})
for k, s in enumerate(plats, start=r0 + 2):
    wn.cell(row=k, column=1, value=s)
    wn.cell(row=k, column=2, value=f'=COUNTIF({IDX_PLT},A{k})')
    wn.cell(row=k, column=3, value=f'=IFERROR(ROUND(AVERAGEIF({IDX_PLT},A{k},{IDX_SC}),1),"")')
    wn.cell(row=k, column=4, value=f'=IFERROR(_xlfn.MAXIFS({IDX_SC},{IDX_PLT},A{k}),"")')
    wn.cell(row=k, column=5, value=f'=IFERROR(_xlfn.MINIFS({IDX_SC},{IDX_PLT},A{k}),"")')
    for col in range(1, 6):
        c = wn.cell(row=k, column=col); c.border = box; c.font = f_base
        if col >= 3: c.number_format = "0.0"
r1 = r0 + 2 + len(plats) + 2
wn.cell(row=r1, column=1, value="Pass rate per check (share of the index scoring 2 / 1 / 0 / U)").font = f_bold
for i, h in enumerate(["Check", "ID", "Pillar", "Meets (2)", "Partly (1)", "Fails (0)", "Unverified", "Meets %", "Failure rank"], 1):
    wn.cell(row=r1 + 1, column=i, value=h)
style_header(wn, r1 + 1, 9)
CHID = f"Checks!$C$2:$C${CHK_LAST + 300}"; CHSC = f"Checks!$E$2:$E${CHK_LAST + 300}"
c_first = r1 + 2
for k, (cid, name, *_) in enumerate(CHECKS, start=c_first):
    wn.cell(row=k, column=1, value=name); wn.cell(row=k, column=2, value=cid); wn.cell(row=k, column=3, value=cid[0])
    wn.cell(row=k, column=4, value=f'=COUNTIFS({CHID},B{k},{CHSC},2)')
    wn.cell(row=k, column=5, value=f'=COUNTIFS({CHID},B{k},{CHSC},1)')
    wn.cell(row=k, column=6, value=f'=COUNTIFS({CHID},B{k},{CHSC},0)')
    wn.cell(row=k, column=7, value=f'=COUNTIFS({CHID},B{k},{CHSC},"U")')
    c = wn.cell(row=k, column=8, value=f'=IFERROR(D{k}/(D{k}+E{k}+F{k}),"")'); c.number_format = "0%"
    wn.cell(row=k, column=9, value=f'=RANK(F{k},$F${c_first}:$F${c_first + len(CHECKS) - 1},0)')
    for col in range(1, 10):
        cc = wn.cell(row=k, column=col); cc.border = box; cc.font = f_base
        if col >= 2: cc.alignment = center
c_last = c_first + len(CHECKS) - 1
wn.conditional_formatting.add(f"H{c_first}:H{c_last}", DataBarRule(start_type="num", start_value=0, end_type="num", end_value=1, color="8A8377"))
wn.conditional_formatting.add(f"F{c_first}:F{c_last}", DataBarRule(start_type="num", start_value=0, end_type="num", end_value=N, color=ACCENT))
wn.cell(row=c_last + 2, column=1, value="Failure rank 1 = the check most stores fail outright. Meets % excludes unverified observations.").font = f_sub

# =====================================================================
# Outreach tracker
# =====================================================================
wo = wb.create_sheet("Outreach")
set_widths(wo, [28, 18, 7, 8, 7, 60, 22, 26, 14, 14, 12, 12, 12, 40])
hdr = ["Company", "Segment", "Rank", "Score", "Grade", "Headline fix (the reason to write)", "Contact", "Title", "Channel", "Status", "Sent", "Follow-up", "Reply", "Notes"]
for i, h in enumerate(hdr, 1):
    wo.cell(row=1, column=i, value=h)
style_header(wo, 1, len(hdr)); wo.row_dimensions[1].height = 30
for k, st in enumerate(rows, start=2):
    wo.cell(row=k, column=1, value=st["company"]).font = f_bold
    wo.cell(row=k, column=2, value=st["segment"])
    wo.cell(row=k, column=3, value=f'=INDEX(Index!$A$2:$A${LAST},MATCH(A{k},Index!$B$2:$B${LAST},0))').font = f_green
    c = wo.cell(row=k, column=4, value=f'=INDEX(Index!$V$2:$V${LAST},MATCH(A{k},Index!$B$2:$B${LAST},0))'); c.font = f_green; c.number_format = "0.0"
    wo.cell(row=k, column=5, value=f'=INDEX(Index!$W$2:$W${LAST},MATCH(A{k},Index!$B$2:$B${LAST},0))').font = f_green
    wo.cell(row=k, column=6, value=st["fixes"][0])
    for col in (7, 8, 9, 11, 12, 13, 14):
        c = wo.cell(row=k, column=col); c.font = f_input; c.fill = fill_input
    s = wo.cell(row=k, column=10, value="Not started"); s.font = f_input; s.fill = fill_input
    for col in range(1, len(hdr) + 1):
        c = wo.cell(row=k, column=col); c.border = box
        if c.font == Font(): c.font = f_base
        c.alignment = Alignment(vertical="top", wrap_text=col in (6, 14), horizontal="center" if col in (3, 4, 5) else None)
    wo.row_dimensions[k].height = 30
# example row showing the expected format
ex = N + 2
wo.cell(row=ex, column=1, value="Example Co (delete me)").font = f_bold
vals = {2: "Retail chain", 6: "Server-render title/meta/canonical tags on product pages", 7: "Jane Doe", 8: "Director, E-commerce", 9: "LinkedIn InMail", 10: "Sent", 11: "2026-10-01", 12: "2026-10-08", 13: "", 14: "Sent the report card + 3-fix note; offered a 20-min walkthrough"}
for col, v in vals.items():
    c = wo.cell(row=ex, column=col, value=v); c.font = f_input if col >= 7 else f_base
for col in range(1, len(hdr) + 1):
    c = wo.cell(row=ex, column=col); c.border = box
    if c.font == Font(): c.font = f_base
    c.alignment = Alignment(vertical="top", wrap_text=col in (6, 14))
dv3 = DataValidation(type="list", formula1='"Not started,Drafted,Sent,Replied,Call booked,Closed"', allow_blank=True)
wo.add_data_validation(dv3); dv3.add(f"J2:J{ex + 100}")
dv4 = DataValidation(type="list", formula1='"LinkedIn InMail,LinkedIn connection note,Email,Phone,Indeed,Referral"', allow_blank=True)
wo.add_data_validation(dv4); dv4.add(f"I2:I{ex + 100}")
wo.freeze_panes = "B2"; wo.auto_filter.ref = f"A1:N{ex}"
wo.cell(row=ex + 2, column=1, value="Fill the yellow cells. Rank, score and grade pull from the Index sheet. Keep the follow-up to one week; log every reply.").font = f_sub

# =====================================================================
# Outreach message
# =====================================================================
wm = wb.create_sheet("Outreach message")
wm.sheet_view.showGridLines = False
set_widths(wm, [3, 110])
wm["B2"] = "Report-card note (send with the company's page from the index)"; wm["B2"].font = f_title
msg = [
    "Subject: Your storefront, scored — [Company] on the ShelfMark index",
    "",
    "Hi [First name],",
    "",
    "I benchmark Michigan storefronts on 29 public-page checks — search visibility, catalog quality, the path to purchase, trust signals and content. [Company] is on the September 2026 index at [Score]/100 ([Grade]), ranked [Rank] of 21. Here is your page: [link]",
    "",
    "Where you're strong: [Strength 1, in one line].",
    "",
    "The one fix I'd make first: [Headline fix — why it matters in one sentence, e.g. 'every product page ships the same meta description, so search results can't tell your products apart'].",
    "",
    "If it would help, I can walk your team through the full report card in 20 minutes — no strings. I'm a webmaster / e-commerce specialist based in Warren, MI (7 years running catalogs of 14,000+ SKUs on Shopify, WooCommerce and Wix), and I'm currently open to full-time roles.",
    "",
    "Moses Batarseh",
    "hireme@mousabatarseh.com · mousabatarseh.com · linkedin.com/in/mousabatarseh",
    "",
    "Rules: keep it under 150 words, name one strength before the fix, never send a fix that isn't on their Findings row, follow up once after 7 days, then move on.",
]
for i, line in enumerate(msg, start=4):
    c = wm.cell(row=i, column=2, value=line); c.font = f_sub if i == 4 + len(msg) - 1 else f_base; c.alignment = wrap
    wm.row_dimensions[i].height = 15 * (len(line) // 105 + 1)

wb.save("ShelfMark-Michigan-Storefront-Benchmark-2026.xlsx")
print("saved", N, "stores; checks rows", CHK_LAST - 1)
